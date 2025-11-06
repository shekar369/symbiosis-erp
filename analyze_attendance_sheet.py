import openpyxl
from openpyxl import load_workbook

# Load the Excel file
file_path = r"C:\Users\Admin\Documents\projects\Claude_exp\HR_Payroll\Doc-refs\1.Input Data\Employee Database.xlsx"
wb = load_workbook(file_path)

# Analyze Attendance Data sheet
ws = wb["Attendance Data"]

print("=== ATTENDANCE DATA SHEET ANALYSIS ===\n")
print(f"Total rows: {ws.max_row}")
print(f"Total columns: {ws.max_column}\n")

# Print first 15 rows to find where actual data starts
print("First 15 rows (showing first 20 columns):")
for row in range(1, min(16, ws.max_row + 1)):
    print(f"\nRow {row}:")
    for col in range(1, min(21, ws.max_column + 1)):
        cell_value = ws.cell(row=row, column=col).value
        if cell_value:  # Only print non-empty cells
            print(f"  Col {col}: {cell_value}")

wb.close()
