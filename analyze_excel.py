import openpyxl
from openpyxl import load_workbook

# Load the Excel file
file_path = r"C:\Users\Admin\Documents\projects\Claude_exp\HR_Payroll\Doc-refs\1.Input Data\Employee Database.xlsx"
wb = load_workbook(file_path)

print("=== EXCEL FILE ANALYSIS ===\n")

# List all sheets
print(f"Sheet names: {wb.sheetnames}\n")

# Analyze each sheet
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\n=== Sheet: {sheet_name} ===")
    print(f"Dimensions: {ws.dimensions}")
    print(f"Max row: {ws.max_row}, Max column: {ws.max_column}")

    # Print headers (first row)
    print("\nHeaders:")
    headers = []
    for col in range(1, ws.max_column + 1):
        cell_value = ws.cell(row=1, column=col).value
        headers.append(cell_value)
        print(f"  Column {col}: {cell_value}")

    # Print first 3 data rows
    print("\nFirst 3 data rows:")
    for row in range(2, min(5, ws.max_row + 1)):
        print(f"\nRow {row}:")
        for col_idx, header in enumerate(headers, start=1):
            cell_value = ws.cell(row=row, column=col_idx).value
            print(f"  {header}: {cell_value}")

wb.close()
