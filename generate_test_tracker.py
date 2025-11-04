#!/usr/bin/env python3
"""
Generate Test Execution Tracking Spreadsheet
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

def create_test_tracker():
    wb = Workbook()

    # Remove default sheet
    if 'Sheet' in wb.sheetnames:
        wb.remove(wb['Sheet'])

    # Define styles
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)

    critical_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
    high_fill = PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid")
    medium_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    low_fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")

    pass_fill = PatternFill(start_color="00FF00", end_color="00FF00", fill_type="solid")
    fail_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
    blocked_fill = PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid")
    pending_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # Test Cases Data
    test_data = {
        "SaaS Admin": [
            ("SA-001", "SaaS Admin Login", "Navigate to login, enter saasadmin/admin123, click Sign In", "Login successful, redirected to dashboard", "Critical"),
            ("SA-002", "Session Persistence", "Login as saasadmin, refresh browser", "User remains logged in", "High"),
            ("SA-003", "Logout", "Login as saasadmin, click logout", "Logged out, redirected to login", "High"),
            ("SA-004", "Invalid Credentials", "Enter wrong password, click Sign In", "Error message displayed", "Medium"),
            ("SA-010", "Dashboard View", "Login as saasadmin, verify dashboard loads", "Dashboard displays with statistics", "Critical"),
            ("SA-011", "Navigation Menu Access", "Check sidebar menu", "All admin menu items visible", "High"),
            ("SA-020", "Access System Config", "Navigate to Admin > System Config", "System config page loads", "High"),
            ("SA-021", "View Configuration Settings", "Check all tabs", "All config tabs display correctly", "Medium"),
            ("SA-022", "Update System Settings", "Modify a setting, click Save", "Settings updated successfully", "High"),
            ("SA-030", "View All Tenants", "Navigate to tenants section", "List of all tenants displayed", "High"),
            ("SA-031", "Access Cross-Tenant Data", "Switch between tenants", "Can view data from multiple tenants", "Critical"),
        ],
        "Employer Admin": [
            ("EA-001", "Employer Admin Login", "Login with employer/employer123", "Login successful, redirected to dashboard", "Critical"),
            ("EA-002", "Session Persistence", "Login, refresh browser", "User remains logged in", "High"),
            ("EA-003", "Logout", "Click logout", "Logged out successfully", "High"),
            ("EA-004", "Tenant Data Isolation", "View employees list", "Only sees own tenant employees", "Critical"),
            ("EA-010", "Dashboard Statistics", "View dashboard", "Shows total, active, exited, attendance %", "Critical"),
            ("EA-011", "Location Filter", "Select location from dropdown", "Statistics update for location", "Medium"),
            ("EA-020", "View Employees List", "Navigate to Employees", "List displayed with filters", "Critical"),
            ("EA-021", "Search Employee", "Enter name in search", "Matching employees shown", "High"),
            ("EA-022", "Filter Employees", "Apply filters (status, location, dept)", "Filtered results displayed", "High"),
            ("EA-023", "Add New Employee", "Click Add, fill fields, save", "Employee created successfully", "Critical"),
            ("EA-024", "View Employee Details", "Click on employee", "Details page displayed", "High"),
            ("EA-025", "Edit Employee", "Click Edit, modify, save", "Employee updated", "High"),
            ("EA-026", "Delete Employee", "Select employee, delete, confirm", "Employee deleted/archived", "High"),
            ("EA-027", "Bulk Upload Employees", "Click Upload, select Excel, upload", "Employees imported from Excel", "High"),
            ("EA-028", "Export Employees", "Click Export, select format", "Employee data exported", "Medium"),
            ("EA-030", "View Locations", "Navigate to Locations", "List of locations displayed", "High"),
            ("EA-031", "Add New Location", "Click Add, fill details, save", "Location created", "High"),
            ("EA-032", "Edit Location", "Click Edit, modify, save", "Location updated", "Medium"),
            ("EA-033", "Delete Location", "Select location, delete", "Location deleted if no employees", "Medium"),
            ("EA-040", "View Attendance", "Navigate to Attendance", "Attendance records with filters", "Critical"),
            ("EA-041", "Filter by Date Range", "Select dates, apply filter", "Attendance for period shown", "High"),
            ("EA-042", "Filter by Location", "Select location, apply", "Attendance for location shown", "High"),
            ("EA-043", "Mark Attendance", "Select date, mark present/absent", "Attendance marked successfully", "Critical"),
            ("EA-044", "Bulk Upload Attendance", "Upload Excel file", "Attendance imported", "High"),
            ("EA-045", "Export Attendance", "Select range, export", "Attendance exported", "Medium"),
            ("EA-050", "View Wage Structure", "Navigate to Wages", "Wage components displayed", "High"),
            ("EA-051", "Define Wage Components", "Add component, save", "Component created", "High"),
            ("EA-052", "Assign Wages to Employee", "Select employee, assign structure", "Wages assigned", "Critical"),
            ("EA-060", "View Leave Requests", "Navigate to Leaves", "Leave requests displayed", "Critical"),
            ("EA-061", "Filter Leave Requests", "Filter by status", "Filtered requests shown", "High"),
            ("EA-062", "Approve Leave Request", "Click on request, approve", "Leave approved", "Critical"),
            ("EA-063", "Reject Leave Request", "Click on request, reject", "Leave rejected with reason", "Critical"),
            ("EA-064", "View Leave Balance", "Select employee, view balance", "Leave balance shown by type", "High"),
            ("EA-070", "View Payroll Dashboard", "Navigate to Payroll", "Payroll dashboard displayed", "Critical"),
            ("EA-071", "Process Monthly Payroll", "Select month, process payroll", "Payroll processed for all", "Critical"),
            ("EA-072", "View Payroll Details", "Select month, click employee", "Detailed salary slip shown", "Critical"),
            ("EA-073", "Apply Deductions", "Select employee, add deduction", "Deduction applied", "High"),
            ("EA-074", "Generate Pay Slips", "Select month, generate", "Pay slips generated", "Critical"),
            ("EA-075", "Download Pay Slips", "Select month, download", "Pay slips downloaded", "High"),
            ("EA-080", "View Statutory Reports", "Navigate to Statutory", "Statutory dashboard shown", "High"),
            ("EA-081", "Generate PF Report", "Select PF, choose dates, generate", "PF report generated", "High"),
            ("EA-082", "Generate ESI Report", "Select ESI, choose dates, generate", "ESI report generated", "High"),
            ("EA-090", "View Bank Transfer", "Navigate to Bank Transfer", "Bank transfer dashboard shown", "High"),
            ("EA-091", "Generate Bank File", "Select month, choose format, generate", "Bank file generated", "Critical"),
            ("EA-092", "Download Bank File", "Generate, download, choose format", "Bank file downloaded", "High"),
            ("EA-100", "View Reports Dashboard", "Navigate to Reports", "Reports dashboard with categories", "High"),
            ("EA-101", "Employee Report", "Select Employee Report, generate", "Employee report generated", "Medium"),
            ("EA-102", "Attendance Report", "Select Attendance Report, generate", "Attendance report generated", "High"),
            ("EA-103", "Payroll Summary Report", "Select Payroll Summary, generate", "Payroll summary generated", "High"),
        ],
        "Employee": [
            ("EM-001", "Employee Login", "Login with employee1/employee123", "Login successful, redirected to dashboard", "Critical"),
            ("EM-002", "Session Persistence", "Login, refresh browser", "User remains logged in", "High"),
            ("EM-003", "Logout", "Click logout", "Logged out successfully", "High"),
            ("EM-004", "Restricted Access", "Try to access employer pages", "Access denied", "Critical"),
            ("EM-010", "View Dashboard", "View dashboard", "Shows personal stats, attendance, leaves", "Critical"),
            ("EM-011", "Quick Actions", "Check quick action buttons", "Shows apply leave, view payslip", "Medium"),
            ("EM-020", "View Profile", "Navigate to Profile", "Complete profile info displayed", "High"),
            ("EM-021", "Update Personal Info", "Click Edit, update fields, save", "Personal info updated", "High"),
            ("EM-022", "Upload Profile Photo", "Click upload, select image, upload", "Profile photo updated", "Low"),
            ("EM-023", "View Employment Details", "View employment tab", "Shows designation, dept, joining date", "Medium"),
            ("EM-024", "View Bank Details", "View bank details section", "Bank info displayed (masked)", "Medium"),
            ("EM-025", "Change Password", "Click Change Password, enter old/new", "Password changed", "High"),
            ("EM-030", "View Attendance", "Navigate to My Attendance", "Personal attendance records shown", "Critical"),
            ("EM-031", "View Attendance Calendar", "View calendar", "Calendar shows present/absent/leave", "Medium"),
            ("EM-033", "Apply for Leave", "Click Apply, select type/dates, submit", "Leave application submitted", "Critical"),
            ("EM-034", "View Leave Balance", "View balance section", "Available leave balance shown", "High"),
            ("EM-035", "View Leave History", "View history tab", "Past leave requests with status", "Medium"),
            ("EM-036", "Cancel Leave Request", "Select pending request, cancel", "Leave request cancelled", "Medium"),
            ("EM-040", "View Payslips", "Navigate to Payslips", "List of payslips by month", "Critical"),
            ("EM-041", "View Payslip Details", "Click on a month", "Detailed payslip shown", "Critical"),
            ("EM-042", "Download Payslip", "Select month, download PDF", "Payslip downloaded as PDF", "Critical"),
            ("EM-043", "View Salary Breakdown", "View payslip breakdown", "Shows earnings, deductions, net pay", "High"),
        ],
        "Cross-Functional": [
            ("CF-001", "Concurrent User Access", "Login multiple users simultaneously", "All users work without conflicts", "High"),
            ("CF-002", "Data Isolation", "Add data as employer, login as employee", "Employee sees only own data", "Critical"),
            ("CF-003", "Leave Approval Flow", "Apply leave as employee, approve as employer", "Leave status updated across roles", "Critical"),
            ("CF-010", "Token Expiration", "Login, wait for token expiry, perform action", "Redirected to login", "Critical"),
            ("CF-011", "SQL Injection Prevention", "Try SQL injection in fields", "Input sanitized, no execution", "Critical"),
            ("CF-012", "XSS Prevention", "Enter script tags in fields", "Scripts escaped, not executed", "Critical"),
            ("CF-013", "CSRF Protection", "Attempt CSRF attack", "Request blocked", "Critical"),
            ("CF-014", "Password Security", "Create user, check database", "Password stored as hash", "Critical"),
            ("CF-020", "Large Dataset Handling", "Upload 1000+ employees, navigate", "Page loads with pagination", "Medium"),
            ("CF-030", "API Failure Handling", "Stop backend, try action", "User-friendly error shown", "High"),
            ("CF-031", "Network Error", "Disconnect internet, try action", "Appropriate error shown", "High"),
            ("CF-032", "Invalid File Upload", "Upload invalid file type", "Error message, file rejected", "Medium"),
        ],
        "Integration": [
            ("INT-001", "API Authentication", "Call protected endpoint without token", "401 Unauthorized returned", "Critical"),
            ("INT-002", "API with Valid Token", "Login, call endpoint with token", "200 OK, data returned", "Critical"),
            ("INT-003", "API Response Format", "Call any endpoint, check response", "Consistent JSON format", "High"),
            ("INT-010", "Data Persistence", "Add employee, restart server, check", "Data persists across restart", "Critical"),
            ("INT-011", "Transaction Rollback", "Perform failing transaction, check DB", "Partial data not saved", "High"),
        ],
        "UI/UX": [
            ("UI-001", "Responsive Design - Desktop", "Open on desktop (1920x1080)", "UI displays properly", "High"),
            ("UI-002", "Responsive Design - Tablet", "Open on tablet (768x1024)", "UI adapts to tablet", "Medium"),
            ("UI-003", "Responsive Design - Mobile", "Open on mobile (375x667)", "Mobile-optimized layout", "Medium"),
            ("UI-004", "Browser - Chrome", "Test all features in Chrome", "All features work", "Critical"),
            ("UI-005", "Browser - Firefox", "Test all features in Firefox", "All features work", "Medium"),
            ("UI-006", "Browser - Edge", "Test all features in Edge", "All features work", "Medium"),
            ("UI-007", "Loading States", "Perform async operation", "Loading indicator shown", "High"),
            ("UI-008", "Success Messages", "Perform successful operation", "Success message displayed", "High"),
            ("UI-009", "Error Messages", "Trigger validation error", "Clear error message shown", "High"),
        ],
        "Regression": [
            ("REG-001", "Application Starts", "Start frontend and backend", "Servers start without errors", "Critical"),
            ("REG-002", "Login Works", "Login with all three roles", "All logins successful", "Critical"),
            ("REG-003", "Dashboard Loads", "View dashboard for each role", "Dashboards display correctly", "Critical"),
            ("REG-004", "Navigation Works", "Navigate between all pages", "Navigation successful", "Critical"),
            ("REG-005", "Logout Works", "Logout for all roles", "Logout successful", "Critical"),
            ("REG-010", "Employee CRUD", "Create, read, update, delete employee", "All operations successful", "Critical"),
            ("REG-011", "Attendance Operations", "Mark and view attendance", "Operations successful", "Critical"),
            ("REG-012", "Leave Operations", "Apply and approve leaves", "Operations successful", "Critical"),
            ("REG-013", "Payroll Processing", "Process monthly payroll", "Payroll processed successfully", "Critical"),
            ("REG-014", "Report Generation", "Generate and export reports", "Reports generated successfully", "Critical"),
        ],
    }

    # Create Summary Sheet
    summary = wb.create_sheet("Summary", 0)
    summary.sheet_properties.tabColor = "1072BA"

    # Summary Header
    summary['A1'] = "HR PAYROLL SYSTEM - TEST EXECUTION SUMMARY"
    summary['A1'].font = Font(bold=True, size=16, color="366092")
    summary.merge_cells('A1:H1')

    summary['A2'] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    summary['A2'].font = Font(italic=True, size=10)
    summary.merge_cells('A2:H2')

    # Test Environment Info
    summary['A4'] = "TEST ENVIRONMENT"
    summary['A4'].font = Font(bold=True, size=12)
    summary['A5'] = "Frontend URL:"
    summary['B5'] = "http://127.0.0.1:5174"
    summary['A6'] = "Backend API:"
    summary['B6'] = "http://127.0.0.1:8000"
    summary['A7'] = "API Docs:"
    summary['B7'] = "http://127.0.0.1:8000/docs"

    # Test Credentials
    summary['A9'] = "TEST CREDENTIALS"
    summary['A9'].font = Font(bold=True, size=12)
    cred_headers = ['Role', 'Username', 'Password', 'Purpose']
    for col, header in enumerate(cred_headers, 1):
        cell = summary.cell(row=10, column=col)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

    credentials = [
        ('SaaS Admin', 'saasadmin', 'admin123', 'System-wide administration'),
        ('Employer Admin', 'employer', 'employer123', 'Company HR management'),
        ('Employee', 'employee1', 'employee123', 'Employee self-service'),
    ]

    for row_idx, cred in enumerate(credentials, 11):
        for col_idx, value in enumerate(cred, 1):
            cell = summary.cell(row=row_idx, column=col_idx)
            cell.value = value
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='left', vertical='center')

    # Test Progress Summary
    summary['A15'] = "TEST EXECUTION PROGRESS"
    summary['A15'].font = Font(bold=True, size=12)

    progress_headers = ['Category', 'Total', 'Passed', 'Failed', 'Blocked', 'Not Tested', 'Pass %']
    for col, header in enumerate(progress_headers, 1):
        cell = summary.cell(row=16, column=col)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

    categories = list(test_data.keys())
    row = 17
    for category in categories:
        total = len(test_data[category])
        summary.cell(row=row, column=1).value = category
        summary.cell(row=row, column=2).value = total
        summary.cell(row=row, column=3).value = 0  # Passed
        summary.cell(row=row, column=4).value = 0  # Failed
        summary.cell(row=row, column=5).value = 0  # Blocked
        summary.cell(row=row, column=6).value = total  # Not Tested
        summary.cell(row=row, column=7).value = "0%"

        for col in range(1, 8):
            summary.cell(row=row, column=col).border = thin_border
            summary.cell(row=row, column=col).alignment = Alignment(horizontal='center')

        row += 1

    # Total row
    total_tests = sum(len(tests) for tests in test_data.values())
    summary.cell(row=row, column=1).value = "TOTAL"
    summary.cell(row=row, column=1).font = Font(bold=True)
    summary.cell(row=row, column=2).value = total_tests
    summary.cell(row=row, column=2).font = Font(bold=True)
    summary.cell(row=row, column=3).value = 0
    summary.cell(row=row, column=4).value = 0
    summary.cell(row=row, column=5).value = 0
    summary.cell(row=row, column=6).value = total_tests
    summary.cell(row=row, column=7).value = "0%"
    summary.cell(row=row, column=7).font = Font(bold=True)

    for col in range(1, 8):
        summary.cell(row=row, column=col).border = thin_border
        summary.cell(row=row, column=col).fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")

    # Priority Summary
    summary['A' + str(row + 2)] = "PRIORITY BREAKDOWN"
    summary['A' + str(row + 2)].font = Font(bold=True, size=12)

    priority_row = row + 3
    priority_headers = ['Priority', 'Count', 'Status']
    for col, header in enumerate(priority_headers, 1):
        cell = summary.cell(row=priority_row, column=col)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

    # Count priorities
    priority_counts = {'Critical': 0, 'High': 0, 'Medium': 0, 'Low': 0}
    for tests in test_data.values():
        for test in tests:
            priority = test[4]
            priority_counts[priority] = priority_counts.get(priority, 0) + 1

    priority_row += 1
    for priority, count in priority_counts.items():
        summary.cell(row=priority_row, column=1).value = priority
        summary.cell(row=priority_row, column=2).value = count
        summary.cell(row=priority_row, column=3).value = "Not Started"

        # Apply priority color
        if priority == "Critical":
            fill = critical_fill
        elif priority == "High":
            fill = high_fill
        elif priority == "Medium":
            fill = medium_fill
        else:
            fill = low_fill

        summary.cell(row=priority_row, column=1).fill = fill

        for col in range(1, 4):
            summary.cell(row=priority_row, column=col).border = thin_border
            summary.cell(row=priority_row, column=col).alignment = Alignment(horizontal='center')

        priority_row += 1

    # Set column widths for summary
    summary.column_dimensions['A'].width = 20
    summary.column_dimensions['B'].width = 20
    summary.column_dimensions['C'].width = 15
    summary.column_dimensions['D'].width = 30
    summary.column_dimensions['E'].width = 12
    summary.column_dimensions['F'].width = 12
    summary.column_dimensions['G'].width = 12
    summary.column_dimensions['H'].width = 12

    # Create individual test case sheets
    for category, tests in test_data.items():
        # Clean sheet name - remove invalid characters
        sheet_name = category.replace(" ", "_").replace("/", "_").replace("\\", "_")
        ws = wb.create_sheet(sheet_name)

        # Set tab color
        if "Admin" in category:
            ws.sheet_properties.tabColor = "FF6B6B"
        elif "Employee" in category:
            ws.sheet_properties.tabColor = "4ECDC4"
        elif "Cross" in category or "Integration" in category:
            ws.sheet_properties.tabColor = "95E1D3"
        elif "UI" in category:
            ws.sheet_properties.tabColor = "F38181"
        else:
            ws.sheet_properties.tabColor = "AA96DA"

        # Title
        ws['A1'] = f"{category} - Test Cases"
        ws['A1'].font = Font(bold=True, size=14)
        ws.merge_cells('A1:J1')

        # Headers
        headers = ['TC ID', 'Test Case', 'Test Steps', 'Expected Result', 'Priority', 'Status', 'Tester', 'Date', 'Comments', 'Bug ID']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = thin_border

        # Data rows
        for row_idx, test in enumerate(tests, 4):
            ws.cell(row=row_idx, column=1).value = test[0]  # TC ID
            ws.cell(row=row_idx, column=2).value = test[1]  # Test Case
            ws.cell(row=row_idx, column=3).value = test[2]  # Steps
            ws.cell(row=row_idx, column=4).value = test[3]  # Expected
            ws.cell(row=row_idx, column=5).value = test[4]  # Priority
            ws.cell(row=row_idx, column=6).value = "Not Tested"  # Status
            ws.cell(row=row_idx, column=7).value = ""  # Tester
            ws.cell(row=row_idx, column=8).value = ""  # Date
            ws.cell(row=row_idx, column=9).value = ""  # Comments
            ws.cell(row=row_idx, column=10).value = ""  # Bug ID

            # Apply priority color to priority column
            priority = test[4]
            if priority == "Critical":
                ws.cell(row=row_idx, column=5).fill = critical_fill
            elif priority == "High":
                ws.cell(row=row_idx, column=5).fill = high_fill
            elif priority == "Medium":
                ws.cell(row=row_idx, column=5).fill = medium_fill
            else:
                ws.cell(row=row_idx, column=5).fill = low_fill

            # Apply borders and alignment
            for col in range(1, 11):
                cell = ws.cell(row=row_idx, column=col)
                cell.border = thin_border
                cell.alignment = Alignment(vertical='top', wrap_text=True)

        # Set column widths
        ws.column_dimensions['A'].width = 10  # TC ID
        ws.column_dimensions['B'].width = 25  # Test Case
        ws.column_dimensions['C'].width = 40  # Steps
        ws.column_dimensions['D'].width = 35  # Expected
        ws.column_dimensions['E'].width = 12  # Priority
        ws.column_dimensions['F'].width = 12  # Status
        ws.column_dimensions['G'].width = 15  # Tester
        ws.column_dimensions['H'].width = 12  # Date
        ws.column_dimensions['I'].width = 30  # Comments
        ws.column_dimensions['J'].width = 12  # Bug ID

        # Freeze panes
        ws.freeze_panes = 'A4'

    # Create Bug Tracking Sheet
    bug_sheet = wb.create_sheet("Bug_Tracking")
    bug_sheet.sheet_properties.tabColor = "FF0000"

    bug_sheet['A1'] = "BUG TRACKING LOG"
    bug_sheet['A1'].font = Font(bold=True, size=14)
    bug_sheet.merge_cells('A1:J1')

    bug_headers = ['Bug ID', 'TC ID', 'Summary', 'Description', 'Severity', 'Priority', 'Status', 'Assigned To', 'Date Found', 'Date Fixed']
    for col, header in enumerate(bug_headers, 1):
        cell = bug_sheet.cell(row=3, column=col)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border

    # Set column widths
    bug_sheet.column_dimensions['A'].width = 10
    bug_sheet.column_dimensions['B'].width = 10
    bug_sheet.column_dimensions['C'].width = 30
    bug_sheet.column_dimensions['D'].width = 40
    bug_sheet.column_dimensions['E'].width = 12
    bug_sheet.column_dimensions['F'].width = 12
    bug_sheet.column_dimensions['G'].width = 12
    bug_sheet.column_dimensions['H'].width = 15
    bug_sheet.column_dimensions['I'].width = 12
    bug_sheet.column_dimensions['J'].width = 12

    bug_sheet.freeze_panes = 'A4'

    # Save workbook
    filename = f"HR_Payroll_Test_Tracker_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb.save(filename)
    print(f"[SUCCESS] Test tracker created successfully: {filename}")
    print(f"[INFO] Total test cases: {total_tests}")
    print(f"[INFO] Categories: {len(test_data)}")
    print(f"[INFO] Sheets created: {len(wb.sheetnames)}")

    return filename

if __name__ == "__main__":
    create_test_tracker()
