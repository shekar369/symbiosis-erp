import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from io import BytesIO
from app.models.employee import EmployeeSalaryDetails

def test_download_wage_template(client, auth_headers):
    user, headers = auth_headers(role="hr_manager")
    response = client.get("/api/v1/wages/template/download", headers=headers)
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

def test_upload_wages(client, auth_headers, test_employee, db):
    user, headers = auth_headers(role="hr_manager")
    
    # Create a sample Excel file
    wb = Workbook()
    ws = wb.active
    
    # Add headers (rows 1-4 are skipped, headers at row 5)
    # We need to match the structure expected by ExcelParser.parse_wages_file
    # It expects headers at row 5 (index 4)
    
    ws['A1'] = "WAGES STATEMENT"
    # ... rows 2-4 ...
    
    # Row 5 Headers
    headers_row = [
        'S.No', 'Employee Code', 'Employee Name', 'Basic Salary', 'HRA', 
        'Conveyance', 'Special Allowance', 'Other Allowance',
        'PF Employee', 'PF Employer', 'ESI Employee', 'ESI Employer',
        'PT', 'TDS', 'Net Salary'
    ]
    for col_idx, header in enumerate(headers_row, start=1):
        ws.cell(row=5, column=col_idx, value=header)
    
    # Add data at Row 6
    ws.cell(row=6, column=1, value=1)
    ws.cell(row=6, column=2, value=test_employee.employee_code)
    ws.cell(row=6, column=3, value="John Doe")
    ws.cell(row=6, column=4, value=50000) # Basic
    ws.cell(row=6, column=5, value=20000) # HRA
    ws.cell(row=6, column=6, value=5000)  # Conveyance
    
    # Save to BytesIO
    file_content = BytesIO()
    wb.save(file_content)
    file_content.seek(0)
    
    # Upload
    files = {"file": ("wages.xlsx", file_content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    response = client.post("/api/v1/uploads/wages", headers=headers, files=files)
    
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert data["summary"]["uploaded"] == 1
    assert data["summary"]["failed"] == 0
    
    # Verify DB
    salary = db.query(EmployeeSalaryDetails).filter(EmployeeSalaryDetails.employee_id == test_employee.id).first()
    assert salary is not None
    assert float(salary.basic_salary) == 50000.0
    assert float(salary.hra) == 20000.0
    assert float(salary.conveyance_allowance) == 5000.0
