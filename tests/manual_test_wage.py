import os
import sys
from datetime import date
from io import BytesIO
from openpyxl import Workbook
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

# Add project root to path
sys.path.append(os.getcwd())

from app.db.base import Base
from app.main import app
from app.api.dependencies import get_db
from app.core.security import get_password_hash
from app.models.user import User
from app.models.tenant import Tenant
from app.models.employee import Employee, EmployeeSalaryDetails

# Setup Test DB
SQLALCHEMY_DATABASE_URL = "sqlite:///:memory:"
engine = create_engine(SQLALCHEMY_DATABASE_URL, connect_args={"check_same_thread": False})
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)

def override_get_db():
    try:
        db = TestingSessionLocal()
        yield db
    finally:
        db.close()

app.dependency_overrides[get_db] = override_get_db

def setup_db():
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)
    db = TestingSessionLocal()
    
    # Create Tenant
    tenant = Tenant(
        name="Test Company",
        slug="test-company",
        email="test@company.com",
        phone="1234567890",
        address="Test Address"
    )
    db.add(tenant)
    db.commit()
    db.refresh(tenant)
    
    # Create User
    user = User(
        username="hr_manager",
        email="hr@company.com",
        hashed_password=get_password_hash("password"),
        role="hr_manager",
        is_active=True,
        tenant_id=tenant.id
    )
    db.add(user)
    db.commit()
    
    # Create Employee
    employee = Employee(
        tenant_id=tenant.id,
        employee_code="EMP001",
        first_name="John",
        last_name="Doe",
        email="john.doe@testcompany.com",
        phone="1234567890",
        date_of_birth=date(1990, 1, 1),
        date_of_joining=date(2020, 1, 1)
    )
    db.add(employee)
    db.commit()
    db.refresh(employee)
    
    db.close()
    return tenant, user, employee

def get_auth_headers(client):
    response = client.post(
        "/api/v1/auth/login",
        data={"username": "hr_manager", "password": "password"}
    )
    if response.status_code != 200:
        print(f"Login Failed: {response.status_code} {response.text}")
        raise Exception("Login failed")
    token = response.json()["access_token"]
    return {"Authorization": f"Bearer {token}"}

def test_download_template(client, headers):
    print("Testing Template Download...")
    response = client.get("/api/v1/wages/template/download", headers=headers)
    if response.status_code == 200:
        print("SUCCESS: Template download")
    else:
        print(f"FAILURE: Template download {response.status_code} {response.text}")

def test_upload_wages(client, headers, employee_code):
    print("Testing Wage Upload...")
    
    # Create Excel
    wb = Workbook()
    ws = wb.active
    ws['A1'] = "WAGES STATEMENT"
    
    # Headers at row 5
    headers_row = [
        'S.No', 'Employee Code', 'Employee Name', 'Basic Salary', 'HRA', 
        'Conveyance', 'Special Allowance', 'Other Allowance',
        'PF Employee', 'PF Employer', 'ESI Employee', 'ESI Employer',
        'PT', 'TDS', 'Net Salary'
    ]
    for col_idx, header in enumerate(headers_row, start=1):
        ws.cell(row=5, column=col_idx, value=header)
        
    # Data at row 6
    ws.cell(row=6, column=1, value=1)
    ws.cell(row=6, column=2, value=employee_code)
    ws.cell(row=6, column=3, value="John Doe")
    ws.cell(row=6, column=4, value=50000)
    ws.cell(row=6, column=5, value=20000)
    
    file_content = BytesIO()
    wb.save(file_content)
    file_content.seek(0)
    
    files = {"file": ("wages.xlsx", file_content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    response = client.post("/api/v1/uploads/wages", headers=headers, files=files)
    
    if response.status_code == 200:
        data = response.json()
        if data["success"] and data["summary"]["uploaded"] == 1:
            print("SUCCESS: Wage upload")
        else:
            print(f"FAILURE: Wage upload response {data}")
    else:
        print(f"FAILURE: Wage upload {response.status_code} {response.text}")

def verify_db(employee_code):
    print("Verifying Database...")
    db = TestingSessionLocal()
    employee = db.query(Employee).filter(Employee.employee_code == employee_code).first()
    salary = db.query(EmployeeSalaryDetails).filter(EmployeeSalaryDetails.employee_id == employee.id).first()
    
    if salary and salary.basic_salary == 50000:
        print("SUCCESS: Database verification")
    else:
        print("FAILURE: Database verification")
    db.close()

if __name__ == "__main__":
    try:
        tenant, user, employee = setup_db()
        client = TestClient(app)
        headers = get_auth_headers(client)
        
        test_download_template(client, headers)
        test_upload_wages(client, headers, employee.employee_code)
        verify_db(employee.employee_code)
        
    except Exception as e:
        print(f"ERROR: {e}")
